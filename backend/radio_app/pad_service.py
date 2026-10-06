"""
Service PAD — Logique métier pour :
- Copie des fichiers vers les dossiers jour
- Archivage automatique des diffusions passées
- Historique Excel des archives
"""
import os
import shutil
import re
import csv
from datetime import datetime, date as date_type, timedelta
from django.conf import settings
from django.utils import timezone


# ═══════════════════════════════════════════════════
# MAPPING JOURS → DOSSIERS PAD
# ═══════════════════════════════════════════════════
JOURS_DOSSIERS = {
    0: '1-ALATSINAINY',   # Lundi
    1: '2-TALATA',        # Mardi
    2: '3-ALAROBIA',      # Mercredi
    3: '4-ALAKAMISY',     # Jeudi
    4: '5- ZOMA',         # Vendredi
    5: '6- SABOTSY',      # Samedi
    6: '0- ALAHADY',      # Dimanche
}

DOSSIER_ARCHIVES = 'archives'
FICHIER_HISTORIQUE = 'historique_archives.csv'
FICHIER_EXCEL = 'historique_archives.xlsx'


# ═══════════════════════════════════════════════════
# EN-TÊTES DU FICHIER HISTORIQUE
# ═══════════════════════════════════════════════════
ENTETES_HISTORIQUE = [
    'Date archive',
    'Date diffusion',
    'Heure diffusion',
    'Heure archive',
    'Client',
    'Téléphone',
    'Nom fichier',
    'Fichier archive',
]


# ═══════════════════════════════════════════════════
# NETTOYER UN NOM DE FICHIER
# ═══════════════════════════════════════════════════
def nettoyer_nom(texte):
    """Enlève les caractères interdits dans un nom de fichier Windows."""
    if not texte:
        return 'Inconnu'
    texte = re.sub(r'[<>:"/\\|?*]', '_', str(texte))
    texte = texte.strip().replace(' ', '_')
    return texte[:50] or 'Inconnu'


# ═══════════════════════════════════════════════════
# 1. COPIER UN FICHIER VERS LE DOSSIER DU JOUR
# ═══════════════════════════════════════════════════
def copier_vers_dossier_jour(chemin_source, nom_fichier, date_diffusion):
    """
    Copie le fichier source vers le dossier correspondant
    au jour de la semaine de `date_diffusion`.
    """
    if not os.path.exists(chemin_source):
        raise FileNotFoundError(
            f"Fichier source introuvable : {chemin_source}"
        )

    jour_semaine = date_diffusion.weekday()
    dossier_jour = JOURS_DOSSIERS.get(jour_semaine, '2-TALATA')

    dossier_destination = os.path.join(
        settings.PAD_LOCAL_PATH,
        dossier_jour
    )
    os.makedirs(dossier_destination, exist_ok=True)

    chemin_destination = os.path.join(dossier_destination, nom_fichier)
    shutil.copy2(chemin_source, chemin_destination)

    return f"{dossier_jour}/{nom_fichier}"

# ═══════════════════════════════════════════════════
# 2. ARCHIVER UNE DIFFUSION
# ═══════════════════════════════════════════════════
def archiver_diffusion(prog):
    """
    Archive UNE diffusion (Programmation).
    - Copie le fichier dans archives/ avec un nom daté
    - Supprime la copie du PAD (dossier jour)
    - Ajoute une ligne dans l'historique
    - Met à jour le statut de la Programmation
    - Si toutes les diffusions sont archivées → FichierAudio = 'Archivé'
    """
    from django.utils import timezone as tz

    fichier = prog.fichierAudio

    # Fallback : chercher via la commande
    if not fichier and prog.commande:
        fichier = prog.commande.fichiersAudio.first()

    if not fichier:
        raise ValueError(
            f"Aucun fichier audio pour la programmation #{prog.id}"
        )

    if not fichier.cheminOrdinateur:
        raise FileNotFoundError("cheminOrdinateur manquant")

    if not os.path.exists(fichier.cheminOrdinateur):
        raise FileNotFoundError(
            f"Fichier introuvable : {fichier.cheminOrdinateur}"
        )

    # ═══ 1. Infos client ═══
    commande = prog.commande
    client = commande.client if commande else None

    nom_client = nettoyer_nom(client.nom if client else 'Inconnu')
    tel_client = nettoyer_nom(
        client.telephone if client else '0000000000'
    )

    # ═══ 2. Nom du fichier ═══
    nom_original = os.path.basename(fichier.cheminOrdinateur)
    nom_sans_ext, extension = os.path.splitext(nom_original)
    nom_sans_ext = nettoyer_nom(nom_sans_ext)

    # ═══ 3. Date/heure de DIFFUSION ═══
    date_str = prog.dateDiffusion.strftime('%Y-%m-%d')
    heure_str = prog.heureDiffusion.strftime('%Hh%M')

    # ═══ 4. Nom final archive ═══
    nom_archive = (
        f"{date_str}_{heure_str}_"
        f"{nom_client}_{tel_client}_"
        f"{nom_sans_ext}{extension}"
    )

    # ═══ 5. Dossier archives ═══
    dossier_archives = os.path.join(
        settings.PAD_LOCAL_PATH,
        DOSSIER_ARCHIVES
    )
    os.makedirs(dossier_archives, exist_ok=True)

    chemin_archive = os.path.join(dossier_archives, nom_archive)

    # ═══ 6. Copier vers archives ═══
    shutil.copy2(fichier.cheminOrdinateur, chemin_archive)

    # ═══ 7. Supprimer du PAD (dossier jour) ═══
    jour_semaine = prog.dateDiffusion.weekday()
    dossier_jour = JOURS_DOSSIERS.get(jour_semaine, '2-TALATA')
    chemin_pad = os.path.join(
        settings.PAD_LOCAL_PATH,
        dossier_jour,
        nom_original
    )

    if os.path.exists(chemin_pad):
        try:
            os.remove(chemin_pad)
        except Exception as e:
            print(f"[PAD] Impossible de supprimer {chemin_pad} : {e}")

    # ═══ 8. Mise à jour Programmation ═══
    prog.statut = 'Archivé'
    prog.dateTraitement = tz.now()
    prog.save(update_fields=['statut', 'dateTraitement'])

    # ═══ 9. Historique ═══
    try:
        ajouter_ligne_historique(prog, chemin_archive)
    except Exception as e:
        print(f"[PAD] Erreur historique : {e}")

    # ═══ 10. Vérifier si TOUTES les diffusions sont archivées ═══
    if commande:
        reste = commande.programmations.exclude(
            statut='Archivé'
        ).exists()

        if not reste:
            fichier.statut = 'Archivé'
            fichier.dateArchive = tz.now()
            fichier.save(update_fields=['statut', 'dateArchive'])

    return chemin_archive


# ═══════════════════════════════════════════════════
# 3. NETTOYER LES DIFFUSIONS PASSÉES
# ═══════════════════════════════════════════════════
def nettoyer_diffusions_passees():
    """
    Parcourt TOUTES les Programmation terminées
    et archive CHAQUE diffusion (une par une).

    Retourne un dict : {'archives': [...], 'erreurs': [...]}
    """
    from .models import Programmation

    maintenant = timezone.now()
    resultats = {
        'archives': [],
        'erreurs': [],
    }

    # ═══ Récupérer les programmations non archivées ═══
    programmations = Programmation.objects.filter(
        statut__in=['À venir', 'Envoyé au PAD', 'Transfert en cours']
    ).select_related('commande__client', 'fichierAudio')

    for prog in programmations:
        # ═══ Si pas de fichierAudio, essayer de le trouver via la commande ═══
        fichier = prog.fichierAudio

        if not fichier and prog.commande:
            fichier = prog.commande.fichiersAudio.first()

        if not fichier:
            # Aucun fichier pour cette commande → skip
            continue

        # ═══ Attacher le fichier à la programmation si pas encore fait ═══
        if not prog.fichierAudio:
            prog.fichierAudio = fichier
            prog.save(update_fields=['fichierAudio'])

        # ═══ Calculer dt_fin = diffusion + durée ═══
        dt_diffusion = datetime.combine(
            prog.dateDiffusion,
            prog.heureDiffusion
        )
        dt_diffusion = timezone.make_aware(
            dt_diffusion,
            timezone.get_current_timezone()
        )

        # ═══ Durée depuis le FICHIER AUDIO (en secondes) ═══
        duree_sec = fichier.duree or 0

        dt_fin = dt_diffusion + timedelta(seconds=duree_sec)

        # Si pas encore terminé → skip
        if dt_fin > maintenant:
            continue

        # ═══ Archiver cette diffusion ═══
        try:
            chemin = archiver_diffusion(prog)
            resultats['archives'].append({
                'programmation': prog.id,
                'fichier': (
                    prog.fichierAudio.nomFichier
                    if prog.fichierAudio else '-'
                ),
                'archive': chemin,
            })
        except Exception as e:
            resultats['erreurs'].append({
                'programmation': prog.id,
                'fichier': (
                    prog.fichierAudio.nomFichier
                    if prog.fichierAudio else '-'
                ),
                'erreur': str(e),
            })

    return resultats

# ═══════════════════════════════════════════════════
# 4. ENVOYER LES FICHIERS EN ATTENTE
# ═══════════════════════════════════════════════════
def envoyer_fichiers_en_attente():
    """
    Cherche tous les fichiers avec statut 'En attente' et
    tente de les envoyer vers le PAD si leur diffusion
    tombe dans la semaine en cours.

    Appelé par le scheduler chaque minute.
    """
    from .models import FichierAudio

    resultats = {
        'envoyes': [],
        'erreurs': [],
    }

    fichiers = FichierAudio.objects.filter(
        statut='En attente'
    ).select_related('commande__client')

    for fichier in fichiers:
        try:
            chemin = fichier.envoyerVersPAD()
            if chemin:
                resultats['envoyes'].append({
                    'fichier': fichier.nomFichier,
                    'chemins': chemin,
                })
        except Exception as e:
            resultats['erreurs'].append({
                'fichier': fichier.nomFichier,
                'erreur': str(e),
            })

    return resultats


# ═══════════════════════════════════════════════════
# 4. LISTER LES ARCHIVES
# ═══════════════════════════════════════════════════



def lister_archives():
    """
    Retourne la liste des fichiers dans le dossier archives/.
    (Sans le CSV et le Excel)
    """
    dossier_archives = os.path.join(
        settings.PAD_LOCAL_PATH,
        DOSSIER_ARCHIVES
    )

    if not os.path.isdir(dossier_archives):
        return []

    fichiers = []
    for nom in sorted(os.listdir(dossier_archives), reverse=True):

        # Ignorer les fichiers historique
        if nom in (FICHIER_HISTORIQUE, FICHIER_EXCEL):
            continue

        chemin = os.path.join(dossier_archives, nom)
        if not os.path.isfile(chemin):
            continue

        try:
            stat = os.stat(chemin)
            fichiers.append({
                'nom': nom,
                'chemin': chemin,
                'taille': stat.st_size,
                'dateModif': datetime.fromtimestamp(
                    stat.st_mtime
                ).isoformat(),
            })
        except Exception:
            continue

    return fichiers

def ajouter_ligne_historique(prog, chemin_archive):
    """
    Ajoute une ligne dans historique_archives.csv
    pour UNE diffusion (Programmation).
    + régénère historique_archives.xlsx
    """
    dossier_archives = os.path.join(
        settings.PAD_LOCAL_PATH,
        DOSSIER_ARCHIVES
    )
    os.makedirs(dossier_archives, exist_ok=True)

    chemin_csv = os.path.join(dossier_archives, FICHIER_HISTORIQUE)
    fichier_existe = os.path.exists(chemin_csv)

    # ═══ 1. Infos client ═══
    commande = prog.commande
    client = commande.client if commande else None

    nom_client = client.nom if client else '-'
    telephone_client = client.telephone if client else '-'

    # ═══ 2. Date/heure diffusion ═══
    date_diff = prog.dateDiffusion.strftime('%d/%m/%Y')
    heure_diff = prog.heureDiffusion.strftime('%H:%M')

    # ═══ 3. Date/heure archive = fin de diffusion ═══
    dt_diff = datetime.combine(
        prog.dateDiffusion,
        prog.heureDiffusion
    )

    duree_sec = (
        prog.fichierAudio.duree
        if prog.fichierAudio and prog.fichierAudio.duree
        else 0
    )

    dt_archive = dt_diff + timedelta(seconds=duree_sec)

    date_archive = dt_archive.strftime('%d/%m/%Y')
    heure_archive = dt_archive.strftime('%H:%M')

    # ═══ 4. Nom du fichier ═══
    nom_fichier = (
        prog.fichierAudio.nomFichier
        if prog.fichierAudio else '-'
    )

    # ═══ 5. Écrire dans le CSV (ordre = ENTETES_HISTORIQUE) ═══
    with open(
        chemin_csv, mode='a', newline='', encoding='utf-8-sig'
    ) as f:
        writer = csv.writer(f, delimiter=';')

        if not fichier_existe:
            writer.writerow(ENTETES_HISTORIQUE)

        writer.writerow([
            date_archive,     # 1. Date archive
            date_diff,        # 2. Date diffusion
            heure_diff,       # 3. Heure diffusion
            heure_archive,    # 4. Heure archive
            nom_client,       # 5. Client
            telephone_client, # 6. Téléphone
            nom_fichier,      # 7. Nom fichier
            os.path.basename(chemin_archive),  # 8. Fichier archive
        ])

    # ═══ 6. Régénérer l'Excel ═══
    try:
        exporter_historique_excel()
    except Exception as e:
        print(f"[PAD] Erreur Excel : {e}")

    return chemin_csv


# ═══════════════════════════════════════════════════
# 6. EXPORTER L'HISTORIQUE EN EXCEL
# ═══════════════════════════════════════════════════
def exporter_historique_excel():
    """
    Lit le CSV et génère un fichier Excel .xlsx
    dans le dossier archives/.
    """
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter

    dossier_archives = os.path.join(
        settings.PAD_LOCAL_PATH,
        DOSSIER_ARCHIVES
    )
    chemin_csv = os.path.join(dossier_archives, FICHIER_HISTORIQUE)
    chemin_xlsx = os.path.join(dossier_archives, FICHIER_EXCEL)

    if not os.path.exists(chemin_csv):
        return None

    # ═══ 1. Lire le CSV ═══
    lignes = []
    with open(chemin_csv, mode='r', encoding='utf-8-sig') as f:
        reader = csv.reader(f, delimiter=';')
        for ligne in reader:
            if ligne:
                lignes.append(ligne)

    if not lignes:
        return None

    # ═══ 2. Créer le classeur Excel ═══
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Historique Archives'

    # ═══ 3. Styles ═══
    title_font = Font(
        name='Calibri', size=16, bold=True, color='007A4D'
    )
    header_font = Font(
        name='Calibri', size=11, bold=True, color='FFFFFF'
    )
    header_fill = PatternFill(
        start_color='007A4D', end_color='007A4D', fill_type='solid'
    )
    cell_font = Font(name='Calibri', size=10)
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    # ═══ 4. Titre ═══
    nb_colonnes = len(lignes[0])
    ws.merge_cells(
        start_row=1, start_column=1,
        end_row=1, end_column=nb_colonnes
    )
    ws.cell(row=1, column=1).value = (
        'RADIO TSIRY — Historique des Archives'
    )
    ws.cell(row=1, column=1).font = title_font
    ws.cell(row=1, column=1).alignment = center_align
    ws.row_dimensions[1].height = 30

    # ═══ 5. En-têtes (ligne 3) ═══
    row_start = 3
    for col_idx, valeur in enumerate(lignes[0], start=1):
        cell = ws.cell(row=row_start, column=col_idx, value=valeur)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[row_start].height = 22

    # ═══ 5bis. Trier les données par date + heure archive (desc) ═══
    donnees = lignes[1:]  # sans en-têtes

    def cle_tri(ligne):
        # Structure :
        #   ligne[0] = Date archive (jj/mm/aaaa)
        #   ligne[1] = Date diffusion
        #   ligne[2] = Heure diffusion
        #   ligne[3] = Heure archive
        date = ligne[0] if len(ligne) > 0 else ''
        heure = ligne[3] if len(ligne) > 3 else ''
        try:
            j, m, a = date.split('/')
            return f"{a}-{m}-{j} {heure}"
        except Exception:
            return ''

    donnees = sorted(donnees, key=cle_tri, reverse=True)

    # ═══ 6. Données ═══
    for row_offset, ligne in enumerate(donnees, start=1):
        row_num = row_start + row_offset
        for col_idx, valeur in enumerate(ligne, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=valeur)
            cell.font = cell_font
            cell.border = thin_border
            # Colonnes centrées : Date archive, Date diffusion,
            # Heure diffusion, Heure archive, Téléphone
            if col_idx in (1, 2, 3, 4, 6):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # ═══ 7. Largeur des colonnes ═══
    largeurs = [14, 14, 20, 16, 14, 14, 40, 55]
    for i, largeur in enumerate(largeurs[:nb_colonnes], start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur

    # ═══ 8. Sauvegarder ═══
    try:
        wb.save(chemin_xlsx)
        return chemin_xlsx
    except PermissionError:
        # Le fichier est ouvert dans Excel → ignorer
        print(
            f"[PAD] Impossible d'écrire {chemin_xlsx} "
            f"(fichier ouvert dans Excel ?)"
        )
        return None